"""Manual review artifacts and revision loop for formal Route 3 runs."""

from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import deepcopy
import json
from pathlib import Path
import re
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .generation_models import EntityReference, EvidenceRecord, GeneratedCandidate
from .route3_artifacts import (
    Route3CandidateArtifact,
    Route3CandidateProvenance,
    complete_route3_candidate_rerun,
    complete_route3_topic_classification,
    create_route3_candidate_artifact,
    revise_route3_candidate_artifact,
)
from .route3_circuit import ServiceCircuit
from .route3_ddg import Route3DDGVerifierResultStore
from .route3_ids import route3_record_identity
from .route3_openrouter import bind_route3_allocation_client, bind_route3_allocation_panel
from .route3_run_ledger import atomic_write_json


REVIEW_STATE_VERSION = 3
REVIEW_COLUMNS = (
    "id",
    "question",
    "reference_answer",
    "wikipedia_url",
    "topic",
    "human_edited",
    "delete",
    "edited_question",
    "edited_reference_answer",
    "edit_reason",
)
REVIEW_TOPICS = (
    "Science & technology",
    "Politics",
    "Art",
    "Other",
    "Geography",
    "Sports",
    "Music",
    "TV shows",
    "History",
    "Video games",
)
DELETE_VALUES = ("Yes", "No")
RERUN_REJECTION_REASONS = {
    "search_longtail_verifier_error",
    "second_stage_grading_error",
}
TOPIC_CLASSIFICATION_MODEL = "openai/gpt-4.1-mini"
TOPIC_CLASSIFICATION_MAX_TOKENS = 256
REVIEW_MARKDOWN_CHUNK_SIZE = 50


def create_review_state(
    records: Iterable[dict[str, Any]],
    *,
    segment_fingerprints: dict[str, dict[str, Any]],
    segment_artifact_roots: dict[str, str],
) -> dict[str, Any]:
    """Create a review state from formal accepted records."""
    bundles = [_review_bundle(record) for record in records]
    segment_ids = {str(bundle["artifact"]["identity"]["segment_id"]) for bundle in bundles}
    missing_segments = sorted(segment_ids - set(segment_fingerprints))
    if missing_segments:
        raise ValueError(f"missing segment fingerprints: {', '.join(missing_segments)}")
    missing_roots = sorted(segment_ids - set(segment_artifact_roots))
    if missing_roots:
        raise ValueError(f"missing segment artifact roots: {', '.join(missing_roots)}")
    candidate_ids = [str(bundle["artifact"]["id"]) for bundle in bundles]
    if len(candidate_ids) != len(set(candidate_ids)):
        raise ValueError("duplicate candidate ID in accepted review input")
    return {
        "review_state_version": REVIEW_STATE_VERSION,
        "segment_fingerprints": deepcopy(segment_fingerprints),
        "segment_artifact_roots": deepcopy(segment_artifact_roots),
        "candidates": bundles,
    }


def topic_classification_prompt(question: str, reference_answer: str) -> str:
    """Build the fixed ten-label topic-classification prompt."""
    labels = "\n".join(f"- {topic}" for topic in REVIEW_TOPICS)
    return (
        "Classify this factual question into exactly one allowed topic.\n"
        "Return only the topic label exactly as written, with no punctuation or explanation.\n\n"
        f"Allowed topics:\n{labels}\n\n"
        f"Question: {question}\n"
        f"Reference answer: {reference_answer}\n"
    )


def classify_review_topics(
    state: dict[str, Any],
    *,
    classifier: Callable[[Route3CandidateArtifact, str], dict[str, Any]],
    concurrency_limit: int,
) -> dict[str, Any]:
    """Classify pending accepted revisions and retain complete call audits."""
    if concurrency_limit < 1:
        raise ValueError("topic classification concurrency_limit must be positive")
    updated = deepcopy(state)
    pending: list[tuple[int, Route3CandidateArtifact, str]] = []
    for index, bundle in enumerate(updated["candidates"]):
        artifact = Route3CandidateArtifact.from_dict(bundle["artifact"])
        revision = artifact.current_revision
        if revision.status == "accepted" and not revision.delete and not revision.topic:
            prompt = topic_classification_prompt(
                revision.authoritative_question,
                revision.reference_answer,
            )
            pending.append((index, artifact, prompt))

    audits: dict[int, dict[str, Any]] = {}
    with ThreadPoolExecutor(max_workers=concurrency_limit) as executor:
        future_indexes = {
            executor.submit(classifier, artifact, prompt): index
            for index, artifact, prompt in pending
        }
        for future in as_completed(future_indexes):
            audits[future_indexes[future]] = future.result()

    pending_by_index = {
        index: (artifact, prompt)
        for index, artifact, prompt in pending
    }
    for index in sorted(audits):
        artifact, prompt = pending_by_index[index]
        response = audits[index]
        raw_text = str(response["raw_text"])
        parsed_topic = str(response["text"]).strip()
        valid = parsed_topic in REVIEW_TOPICS
        topic = parsed_topic if valid else ""
        status = "accepted" if valid else "rejected"
        artifact = complete_route3_topic_classification(
            artifact,
            topic=topic,
            status=status,
        )
        bundle = updated["candidates"][index]
        bundle["artifact"] = artifact.to_dict()
        bundle["active_record"]["topic"] = topic
        bundle["topic_classifications"].append(
            {
                "revision_number": artifact.current_revision.revision_number,
                "model": TOPIC_CLASSIFICATION_MODEL,
                "max_tokens": TOPIC_CLASSIFICATION_MAX_TOKENS,
                "prompt": prompt,
                "request_payload": deepcopy(response["request_payload"]),
                "raw_response": deepcopy(response["response_body"]),
                "raw_response_body_text": str(response.get("raw_response_body_text", "")),
                "raw_text": raw_text,
                "parsed_topic": topic,
                "status": status,
                "rejection_reason": "" if valid else "invalid_topic_classification_response",
            }
        )
    return updated


def load_review_state(path: Path) -> dict[str, Any]:
    """Load and validate one review-state JSON document."""
    state = json.loads(path.read_text(encoding="utf-8"))
    if int(state["review_state_version"]) != REVIEW_STATE_VERSION:
        raise ValueError("unsupported review state version")
    for bundle in state["candidates"]:
        Route3CandidateArtifact.from_dict(bundle["artifact"])
    return state


def write_review_state(path: Path, state: dict[str, Any]) -> None:
    """Atomically write one review state."""
    atomic_write_json(path, state)


def accepted_review_bundles(state: dict[str, Any]) -> list[dict[str, Any]]:
    """Return bundles whose latest revision is active and accepted."""
    rows: list[dict[str, Any]] = []
    for bundle in state["candidates"]:
        artifact = Route3CandidateArtifact.from_dict(bundle["artifact"])
        revision = artifact.current_revision
        if revision.status == "accepted" and not revision.delete:
            if revision.topic not in REVIEW_TOPICS:
                raise ValueError(
                    f"topic classification is incomplete: {artifact.candidate_id}"
                )
            rows.append(bundle)
    return sorted(rows, key=lambda bundle: str(bundle["artifact"]["id"]))


def render_review_markdown(
    state: dict[str, Any],
    *,
    run_id: str,
    bundles: list[dict[str, Any]] | None = None,
    start_index: int = 1,
) -> str:
    """Render accepted candidates with source tables and isolated model responses."""
    selected = accepted_review_bundles(state) if bundles is None else bundles
    lines = [
        f"# Route 3 Review: `{run_id}`",
        "",
        f"Candidates in this file: {len(selected)}",
        "",
    ]
    for index, bundle in enumerate(selected, start=start_index):
        artifact = Route3CandidateArtifact.from_dict(bundle["artifact"])
        revision = artifact.current_revision
        provenance = artifact.provenance
        lines.extend(
            [
                f"## {index}. `{artifact.candidate_id}`",
                "",
                f"Question: {revision.authoritative_question}",
                "",
                f"Reference answer: **{revision.reference_answer}**",
                "",
                f"Topic: `{revision.topic}`",
                "",
                f"Human edited: `{_human_edited(artifact)}`",
                "",
                f"Answer type: `{provenance.answer_type}`",
                "",
                f"Wikipedia: [{provenance.canonical_page_url}]({provenance.canonical_page_url})",
                "",
                f"Canonical page ID: `{provenance.canonical_page_id}`",
                "",
                "Selected table:",
                "",
                str(provenance.selected_table["markdown"]),
                "",
                "Two-model answers:",
                "",
            ]
        )
        for model in revision.second_stage["models"]:
            lines.extend(
                [
                    f"**{model['model']}**",
                    "",
                    "Model response:",
                    "",
                    _blockquote(str(model["predicted_answer"])),
                    "",
                    f"Judge result: `{model['grade']}`",
                    "",
                    "Judge reason:",
                    "",
                    _blockquote(str(model["reason"])),
                    "",
                ]
            )
    return "\n".join(lines).rstrip() + "\n"



def write_review_markdown_shards(
    base_path: Path,
    state: dict[str, Any],
    *,
    run_id: str,
) -> list[Path]:
    """Write fixed 50-candidate Markdown shards and return their paths."""
    base_path.parent.mkdir(parents=True, exist_ok=True)
    base_path.unlink(missing_ok=True)
    shard_name = re.compile(
        rf"{re.escape(base_path.stem)}_\d+-\d+{re.escape(base_path.suffix)}"
    )
    for stale_path in base_path.parent.iterdir():
        if shard_name.fullmatch(stale_path.name):
            stale_path.unlink()
    bundles = accepted_review_bundles(state)
    chunks = [
        bundles[index:index + REVIEW_MARKDOWN_CHUNK_SIZE]
        for index in range(0, len(bundles), REVIEW_MARKDOWN_CHUNK_SIZE)
    ] or [[]]
    paths: list[Path] = []
    for chunk_index, chunk in enumerate(chunks):
        start = chunk_index * REVIEW_MARKDOWN_CHUNK_SIZE + 1
        end = start + REVIEW_MARKDOWN_CHUNK_SIZE - 1
        path = base_path.with_name(f"{base_path.stem}_{start}-{end}{base_path.suffix}")
        path.write_text(
            render_review_markdown(
                state,
                run_id=run_id,
                bundles=chunk,
                start_index=start,
            ),
            encoding="utf-8",
        )
        paths.append(path)
    return paths


def _blockquote(value: str) -> str:
    """Render every response line as Markdown blockquote content."""
    return "\n".join(f"> {line}" if line else ">" for line in value.splitlines() or [""])


def _human_edited(artifact: Route3CandidateArtifact) -> str:
    """Return whether any authoritative Q/A revision changed from its predecessor."""
    previous = artifact.revisions[0]
    for revision in artifact.revisions[1:]:
        if (
            revision.authoritative_question != previous.authoritative_question
            or revision.reference_answer != previous.reference_answer
        ):
            return "Yes"
        previous = revision
    return "No"

def write_review_workbook(path: Path, state: dict[str, Any]) -> None:
    """Write the formal English-column review workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "review"
    worksheet.append(list(REVIEW_COLUMNS))
    for bundle in accepted_review_bundles(state):
        artifact = Route3CandidateArtifact.from_dict(bundle["artifact"])
        revision = artifact.current_revision
        worksheet.append(
            [
                artifact.candidate_id,
                revision.authoritative_question,
                revision.reference_answer,
                artifact.provenance.canonical_page_url,
                revision.topic,
                "No",
                "No",
                "",
                "",
                "",
            ]
        )
    last_row = max(2, worksheet.max_row)
    delete_validation = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    worksheet.add_data_validation(delete_validation)
    delete_validation.add(f"G2:G{last_row}")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = (42, 58, 28, 48, 24, 15, 12, 58, 30, 42)
    for column_index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def read_review_workbook(path: Path) -> list[dict[str, str]]:
    """Read formal review rows and require the exact English header order."""
    workbook = load_workbook(path, data_only=False)
    worksheet = workbook["review"]
    values = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    headers = tuple(str(value or "") for value in values[0])
    if headers != REVIEW_COLUMNS:
        raise ValueError(f"Review columns must exactly match: {', '.join(REVIEW_COLUMNS)}")
    rows: list[dict[str, str]] = []
    for values_row in values[1:]:
        if all(value is None or str(value).strip() == "" for value in values_row):
            continue
        rows.append(
            {
                column: str(values_row[index] or "").strip()
                for index, column in enumerate(REVIEW_COLUMNS)
            }
        )
    return rows


def validate_review_rows(
    rows: Iterable[dict[str, str]],
    *,
    known_ids: set[str],
    finalization: bool = False,
) -> None:
    """Validate IDs, topics, deletion values, and edit combinations."""
    seen_ids: set[str] = set()
    for row in rows:
        candidate_id = row["id"]
        if candidate_id in seen_ids:
            raise ValueError(f"duplicate review ID: {candidate_id}")
        if candidate_id not in known_ids:
            raise ValueError(f"unknown review ID: {candidate_id}")
        seen_ids.add(candidate_id)
        topic = row["topic"]
        if topic and topic not in REVIEW_TOPICS:
            raise ValueError(f"invalid review topic: {topic}")
        if finalization and not topic:
            raise ValueError(f"topic is required for finalization: {candidate_id}")
        delete_value = row["delete"]
        if row["human_edited"] not in DELETE_VALUES:
            raise ValueError(f"invalid human_edited value for {candidate_id}: {row['human_edited']}")
        if delete_value not in DELETE_VALUES:
            raise ValueError(f"invalid delete value for {candidate_id}: {delete_value}")
        if delete_value == "Yes" and (row["edited_question"] or row["edited_reference_answer"]):
            raise ValueError(f"deleted row cannot contain edited Q/A: {candidate_id}")


def apply_review_rows(state: dict[str, Any], rows: Iterable[dict[str, str]]) -> dict[str, Any]:
    """Apply review edits as append-only revisions without running network checks."""
    updated = deepcopy(state)
    active = accepted_review_bundles(updated)
    known_ids = {str(bundle["artifact"]["id"]) for bundle in active}
    row_list = list(rows)
    validate_review_rows(row_list, known_ids=known_ids)
    rows_by_id = {row["id"]: row for row in row_list}
    for bundle in updated["candidates"]:
        artifact = Route3CandidateArtifact.from_dict(bundle["artifact"])
        row = rows_by_id.get(artifact.candidate_id)
        if row is None:
            continue
        current = artifact.current_revision
        if row["question"] != current.authoritative_question:
            raise ValueError(f"XLSX question does not match candidate revision: {artifact.candidate_id}")
        if row["reference_answer"] != current.reference_answer:
            raise ValueError(f"XLSX answer does not match candidate revision: {artifact.candidate_id}")
        if row["wikipedia_url"] != artifact.provenance.canonical_page_url:
            raise ValueError(f"XLSX URL does not match candidate provenance: {artifact.candidate_id}")
        if row["topic"] != current.topic:
            raise ValueError(f"XLSX topic is read-only: {artifact.candidate_id}")
        if row["human_edited"] != _human_edited(artifact):
            raise ValueError(f"XLSX human_edited is read-only: {artifact.candidate_id}")
        edited_question = row["edited_question"] or None
        edited_answer = row["edited_reference_answer"] or None
        question_changed = (
            edited_question is not None
            and edited_question != current.authoritative_question
        )
        answer_changed = (
            edited_answer is not None
            and edited_answer != current.reference_answer
        )
        qa_changed = question_changed or answer_changed
        delete = row["delete"] == "Yes"
        if not qa_changed and delete == current.delete:
            continue
        artifact = revise_route3_candidate_artifact(
            artifact,
            edited_question=edited_question if question_changed else None,
            edited_reference_answer=edited_answer if answer_changed else None,
            topic="" if qa_changed else None,
            delete=delete,
            edit_reason=row["edit_reason"],
        )
        bundle["artifact"] = artifact.to_dict()
        active_record = bundle["active_record"]
        active_record["topic"] = artifact.current_revision.topic
        if qa_changed:
            active_record["question"] = artifact.current_revision.authoritative_question
            active_record["canonical_question"] = artifact.current_revision.authoritative_question
            active_record["answer"] = artifact.current_revision.reference_answer
            active_record["answer_aliases"] = list(artifact.current_revision.active_aliases)
            active_record["search_queries"] = list(artifact.current_revision.active_search_queries)
            active_record["validation"] = {}
            active_record["search_verification_features"] = {}
            active_record["panel_grading_features"] = {}
            active_record["source_metadata"].pop("rule_based_qa_gate", None)
    return updated


def rerun_review_candidates(
    state: dict[str, Any],
    *,
    processor: Callable[[GeneratedCandidate], tuple[str, dict[str, Any]]],
) -> dict[str, Any]:
    """Run fresh post-generation processing for every active rerun revision."""
    updated = deepcopy(state)
    for bundle in updated["candidates"]:
        artifact = Route3CandidateArtifact.from_dict(bundle["artifact"])
        if artifact.current_revision.status != "rerun":
            continue
        candidate = _generated_candidate(bundle["active_record"], artifact)
        status, processed_record = processor(candidate)
        processed_record["id"] = artifact.candidate_id
        processed_record["topic"] = artifact.current_revision.topic
        metadata = processed_record["source_metadata"]
        artifact = complete_route3_candidate_rerun(
            artifact,
            source_validation=processed_record["validation"],
            integrated_answer_type_gate=metadata.get("rule_based_qa_gate", {}),
            ddg=processed_record["search_verification_features"],
            second_stage=processed_record["panel_grading_features"],
            status=status,
        )
        bundle["artifact"] = artifact.to_dict()
        bundle["active_record"] = processed_record
    return updated


def post_generation_processor(
    *,
    settings,
    search_client,
    second_stage_model_clients,
    grading_grader_client,
    external_call_record_root: Path,
    ddg_verifier_result_root: Path,
    segment_fingerprint: str,
    duckduckgo_circuit: ServiceCircuit,
) -> Callable[[GeneratedCandidate], tuple[str, dict[str, Any]]]:
    """Build a processor that executes the formal post-generation pipeline."""
    from .generation_pipeline import process_generated_candidates

    ddg_verifier_result_store = Route3DDGVerifierResultStore(
        root=ddg_verifier_result_root,
        segment_fingerprint=segment_fingerprint,
        circuit=duckduckgo_circuit,
    )

    def process(candidate: GeneratedCandidate) -> tuple[str, dict[str, Any]]:
        canonical_page_id = int(candidate.source_metadata["canonical_page_id"])
        bound_panel = bind_route3_allocation_panel(
            second_stage_model_clients,
            record_root=external_call_record_root,
            canonical_page_id=canonical_page_id,
        )
        bound_grader = (
            bind_route3_allocation_client(
                grading_grader_client,
                record_root=external_call_record_root,
                canonical_page_id=canonical_page_id,
                call_key="",
            )
            if grading_grader_client is not None
            else None
        )
        result = process_generated_candidates(
            [candidate],
            settings=settings,
            search_client=search_client,
            ddg_verifier_result_store=ddg_verifier_result_store,
            rewrite_client=None,
            second_stage_model_clients=bound_panel,
            grading_grader_client=bound_grader,
        )
        if result.accepted:
            return "accepted", result.accepted[0]
        record = result.rejected[0]
        status = "rerun" if str(record["rejection_reason"]) in RERUN_REJECTION_REASONS else "rejected"
        return status, record

    return process


def _review_bundle(record: dict[str, Any]) -> dict[str, Any]:
    metadata = record["source_metadata"]
    identity = route3_record_identity(record)
    llm_audit = metadata["llm_audit"]
    selected_table = metadata["selected_source_table"]
    provenance = Route3CandidateProvenance(
        run_group_id=identity.run_group_id,
        segment_id=identity.segment_id,
        page_attempt=int(metadata["page_attempt"]),
        canonical_page_id=identity.canonical_page_id,
        canonical_page_url=str(metadata["canonical_url"]),
        selected_table=deepcopy(selected_table),
        selected_table_type=str(selected_table["table_type"]),
        page_archive_hash=str(metadata["route3_page_archive"]["archive_sha256"]),
        generation_prompt=str(metadata["llm_prompt"]),
        generation_request=deepcopy(llm_audit["request_payload"]),
        generation_raw_response=deepcopy(metadata["llm_response"]),
        original_question=str(record["question"]),
        original_reference_answer=str(record["answer"]),
        original_aliases=tuple(str(value) for value in record["answer_aliases"]),
        original_search_queries=tuple(str(value) for value in record["search_queries"]),
        answer_type=str(record["answer_type"]),
        generation_model=str(metadata["generation_model"]),
        generation_parameters=deepcopy(metadata["generation_parameters"]),
        recipe_seed=metadata["recipe_seed"],
    )
    artifact = create_route3_candidate_artifact(
        identity=identity,
        provenance=provenance,
        question=str(record["question"]),
        reference_answer=str(record["answer"]),
        aliases=tuple(str(value) for value in record["answer_aliases"]),
        search_queries=tuple(str(value) for value in record["search_queries"]),
        topic="",
        source_validation=record["validation"],
        integrated_answer_type_gate=metadata["rule_based_qa_gate"],
        ddg=record["search_verification_features"],
        second_stage=record["panel_grading_features"],
        status="accepted",
    )
    return {
        "artifact": artifact.to_dict(),
        "active_record": deepcopy(record),
        "topic_classifications": [],
    }


def _generated_candidate(record: dict[str, Any], artifact: Route3CandidateArtifact) -> GeneratedCandidate:
    revision = artifact.current_revision
    subject = record["subject_entity"]
    evidence = record["evidence"]
    source_metadata = deepcopy(record["source_metadata"])
    source_metadata["route3_revision_number"] = revision.revision_number
    return GeneratedCandidate(
        source_type=str(record["source_type"]),
        generation_route=str(record["generation_route"]),
        question=revision.authoritative_question,
        answer=revision.reference_answer,
        answer_aliases=list(revision.active_aliases),
        subject_entity=EntityReference(**subject),
        answer_entity=EntityReference(name=revision.reference_answer),
        relation_or_claim=str(record["relation_or_claim"]),
        evidence=EvidenceRecord(**evidence),
        canonical_question=revision.authoritative_question,
        question_family=str(record["question_family"]),
        answer_type=artifact.provenance.answer_type,
        topic=revision.topic,
        target_time=str(record["target_time"]),
        source_template_domain=str(record["template_key"]),
        search_queries=list(revision.active_search_queries),
        source_metadata=source_metadata,
    )
