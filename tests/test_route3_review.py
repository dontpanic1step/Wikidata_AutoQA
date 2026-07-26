"""Tests for the formal Route 3 manual review loop."""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from test_support import ROOT  # noqa: F401
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from run_route3_review import main as review_main  # noqa: E402
from wikidata_simpleqa.route3_artifacts import Route3CandidateArtifact
from wikidata_simpleqa.route3_ids import route3_record_id
from wikidata_simpleqa.route3_review import (
    REVIEW_COLUMNS,
    REVIEW_TOPICS,
    accepted_review_bundles,
    apply_review_rows,
    classify_review_topics,
    create_review_state as _create_review_state,
    read_review_workbook,
    render_review_markdown,
    rerun_review_candidates,
    validate_review_rows,
    write_review_workbook,
    write_review_markdown_shards,
)


def accepted_record(
    *,
    page_id: int = 123,
    slot: str = "single",
    answer_type: str = "Person",
    segment_id: str = "01_alltypes_10",
) -> dict:
    """Build one complete formal accepted record."""
    record = {
        "id": "",
        "question": "Who founded the archive?",
        "canonical_question": "Who founded the archive?",
        "answer": "Ada Example",
        "answer_aliases": ["A. Example"],
        "source_type": "wikipedia_tables",
        "generation_route": "route3_wikipedia_infobox",
        "subject_entity": {
            "name": "Archive",
            "qid": "",
            "wikipedia_title": "Archive",
            "url": "https://en.wikipedia.org/wiki/Archive",
        },
        "answer_entity": {
            "name": "Ada Example",
            "qid": "",
            "wikipedia_title": "",
            "url": "",
        },
        "relation_or_claim": "founder",
        "evidence": {
            "text": "| Founder | Ada Example |",
            "url": "https://en.wikipedia.org/wiki/Archive",
            "source_title": "Archive",
            "section": "",
            "retrieved_at": "2026-07-24T00:00:00Z",
        },
        "question_family": "wikipedia_table_single_fact",
        "answer_type": answer_type,
        "template_key": "wikipedia_infobox_table",
        "topic": "",
        "target_time": "2024",
        "search_queries": ["archive founder"],
        "search_verification_features": {
            "passed": True,
            "category_hit_rates": {"overall": {"answer_hit_rate": 0.0}},
        },
        "panel_grading_features": {
            "models": [
                {
                    "model": "openai/gpt-4.1-mini",
                    "predicted_answer": "Ada Example",
                    "grade": "CORRECT",
                    "reason": "Matched.",
                },
                {
                    "model": "google/gemini-3-flash-preview",
                    "predicted_answer": "Ada Example",
                    "grade": "CORRECT",
                    "reason": "Matched.",
                },
            ]
        },
        "validation": {"passed": True},
        "source_metadata": {
            "run_group_id": "review-run",
            "segment_id": segment_id,
            "canonical_page_id": page_id,
            "original_candidate_slot": slot,
            "page_attempt": 1,
            "canonical_url": "https://en.wikipedia.org/wiki/Archive",
            "selected_source_table": {
                "table_type": "infobox",
                "markdown": "| Field | Value |\n| --- | --- |\n| Founder | Ada Example |",
            },
            "route3_page_archive": {"archive_sha256": "abc123"},
            "llm_prompt": "Generate one question.",
            "llm_response": {"question": "Who founded the archive?"},
            "llm_audit": {"request_payload": {"messages": []}},
            "generation_model": "google/gemini-3-flash-preview",
            "generation_parameters": {"max_tokens": 4096},
            "recipe_seed": 42,
            "rule_based_qa_gate": {"status": "matched"},
        },
    }
    record["id"] = route3_record_id(record)
    return record


def review_row(candidate_id: str, **overrides: str) -> dict[str, str]:
    """Build one valid workbook row."""
    row = {
        "id": candidate_id,
        "question": "Who founded the archive?",
        "reference_answer": "Ada Example",
        "wikipedia_url": "https://en.wikipedia.org/wiki/Archive",
        "human_edited": "No",
        "topic": "History",
        "delete": "No",
        "edited_question": "",
        "edited_reference_answer": "",
        "edit_reason": "",
    }
    row.update(overrides)
    return row


def fake_topic_classifier(_artifact, prompt: str) -> dict:
    """Return one complete deterministic topic-classification audit."""
    return {
        "text": "History",
        "raw_text": "History",
        "request_payload": {
            "model": "openai/gpt-4.1-mini",
            "max_tokens": 256,
            "messages": [{"role": "user", "content": prompt}],
        },
        "response_body": {"id": "topic-response", "choices": []},
        "raw_response_body_text": '{"id":"topic-response","choices":[]}',
    }


def create_review_state(*args, **kwargs) -> dict:
    """Create a deterministically topic-classified test review state."""
    state = _create_review_state(*args, **kwargs)
    return classify_review_topics(
        state,
        classifier=fake_topic_classifier,
        concurrency_limit=2,
    )


def test_review_state_requires_each_top_up_segment_fingerprint() -> None:
    records = [
        accepted_record(page_id=123, segment_id="01_alltypes_10"),
        accepted_record(page_id=456, segment_id="01_alltypes_10_topup_01"),
    ]

    with pytest.raises(ValueError, match="missing segment fingerprints"):
        create_review_state(
            records,
            segment_fingerprints={"01_alltypes_10": {"sha256": "first"}},
            segment_artifact_roots={
                "01_alltypes_10": "C:/artifacts/01_alltypes_10",
                "01_alltypes_10_topup_01": "C:/artifacts/01_alltypes_10_topup_01",
            },
        )

    state = create_review_state(
        records,
        segment_fingerprints={
            "01_alltypes_10": {"sha256": "first"},
            "01_alltypes_10_topup_01": {"sha256": "second"},
        },
        segment_artifact_roots={
            "01_alltypes_10": "C:/artifacts/01_alltypes_10",
            "01_alltypes_10_topup_01": "C:/artifacts/01_alltypes_10_topup_01",
        },
    )
    assert set(state["segment_fingerprints"]) == {
        "01_alltypes_10",
        "01_alltypes_10_topup_01",
    }
    assert set(state["segment_artifact_roots"]) == {
        "01_alltypes_10",
        "01_alltypes_10_topup_01",
    }

def test_export_cli_writes_review_state_markdown_and_xlsx() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        accepted_path = root / "accepted.jsonl"
        manifest_path = root / "segment_manifest.json"
        state_path = root / "review_state.json"
        markdown_path = root / "review.md"
        workbook_path = root / "review.xlsx"
        accepted_path.write_text(
            json.dumps(accepted_record() | {"topic": "Wikipedia semi-structured data"}) + "\n",
            encoding="utf-8",
        )
        manifest_path.write_text(json.dumps({"segment_id": "01_alltypes_10", "fingerprint": {"sha256": "test"}}), encoding="utf-8")

        assert review_main(
            [
                "export",
                "--accepted-input",
                str(accepted_path),
                "--segment-manifest",
                str(manifest_path),
                "--state-output",
                str(state_path),
                "--markdown-output",
                str(markdown_path),
                "--xlsx-output",
                str(workbook_path),
                "--run-id",
                "review-run",
            ],
            topic_classifier=fake_topic_classifier,
        ) == 0

        assert state_path.exists()
        exported_state = json.loads(state_path.read_text(encoding="utf-8"))
        assert exported_state["review_state_version"] == 3
        assert exported_state["segment_artifact_roots"] == {
            "01_alltypes_10": str(root.resolve())
        }
        artifact = Route3CandidateArtifact.from_dict(exported_state["candidates"][0]["artifact"])
        assert artifact.current_revision.topic == "History"
        assert exported_state["candidates"][0]["topic_classifications"][0][
            "raw_response_body_text"
        ] == '{"id":"topic-response","choices":[]}'
        assert (root / "review_1-50.md").exists()
        assert workbook_path.exists()
        exported_workbook = load_workbook(workbook_path)
        assert exported_workbook["review"]["E2"].value == "History"
        assert exported_workbook["review"]["F2"].value == "No"
        assert exported_workbook["review"]["G2"].value == "No"
        exported_workbook.close()
        workbook = load_workbook(workbook_path)
        workbook["review"]["E2"] = "History"
        workbook.save(workbook_path)
        next_state_path = root / "review_state_next.json"
        next_markdown_path = root / "review_next.md"
        next_workbook_path = root / "review_next.xlsx"
        subprocess.run(
            [
                sys.executable,
                str(ROOT / "scripts" / "run_route3_review.py"),
                "apply",
                "--state-input",
                str(state_path),
                "--xlsx-input",
                str(workbook_path),
                "--state-output",
                str(next_state_path),
                "--markdown-output",
                str(next_markdown_path),
                "--xlsx-output",
                str(next_workbook_path),
                "--run-id",
                "review-run",
            ],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        next_state = json.loads(next_state_path.read_text(encoding="utf-8"))
        next_revision = next_state["candidates"][0]["artifact"]["revisions"][-1]
        assert next_revision["topic"] == "History"
        assert next_revision["status"] == "accepted"
        assert load_workbook(next_workbook_path)["review"]["E2"].value == "History"

def test_markdown_and_workbook_show_only_formal_review_fields() -> None:
    state = create_review_state(
        [accepted_record()],
        segment_fingerprints={"01_alltypes_10": {"sha256": "fingerprint"}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )
    candidate_id = state["candidates"][0]["artifact"]["id"]

    markdown = render_review_markdown(state, run_id="review-run")

    assert candidate_id in markdown
    assert "Who founded the archive?" in markdown
    assert "Ada Example" in markdown
    assert "Canonical page ID: `123`" in markdown
    assert "| Founder | Ada Example |" in markdown
    assert "openai/gpt-4.1-mini" in markdown
    assert "google/gemini-3-flash-preview" in markdown
    assert "Human edited: `No`" in markdown
    assert "> Ada Example" in markdown
    assert "Judge result: `CORRECT`" in markdown

    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "review.xlsx"
        write_review_workbook(path, state)
        workbook = load_workbook(path)
        worksheet = workbook["review"]
        assert tuple(cell.value for cell in worksheet[1]) == REVIEW_COLUMNS
        assert worksheet["F2"].value == "No"
        assert worksheet["G2"].value == "No"
        assert len(worksheet.data_validations.dataValidation) == 1
        formulas = {validation.formula1 for validation in worksheet.data_validations.dataValidation}
        assert read_review_workbook(path)[0]["id"] == candidate_id
        assert formulas == {'"Yes,No"'}


def test_workbook_rejects_noncanonical_header() -> None:
    state = create_review_state(
        [accepted_record()],
        segment_fingerprints={"01_alltypes_10": {}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )
    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "review.xlsx"
        write_review_workbook(path, state)
        workbook = load_workbook(path)
        workbook["review"]["A1"] = "编号"
        workbook.save(path)

        with pytest.raises(ValueError, match="exactly match"):
            read_review_workbook(path)


def test_review_validation_rejects_invalid_rows() -> None:
    candidate_id = accepted_record()["id"]
    known_ids = {candidate_id}

    with pytest.raises(ValueError, match="duplicate"):
        validate_review_rows(
            [review_row(candidate_id), review_row(candidate_id)],
            known_ids=known_ids,
        )
    with pytest.raises(ValueError, match="unknown"):
        validate_review_rows([review_row("unknown")], known_ids=known_ids)
    with pytest.raises(ValueError, match="invalid review topic"):
        validate_review_rows([review_row(candidate_id, topic="Invalid")], known_ids=known_ids)
    with pytest.raises(ValueError, match="topic is required"):
        validate_review_rows(
            [review_row(candidate_id, topic="")],
            known_ids=known_ids,
            finalization=True,
        )
    with pytest.raises(ValueError, match="invalid delete"):
        validate_review_rows([review_row(candidate_id, delete="Maybe")], known_ids=known_ids)
    with pytest.raises(ValueError, match="deleted row"):
        validate_review_rows(
            [review_row(candidate_id, delete="Yes", edited_question="Edited?")],
            known_ids=known_ids,
        )


def test_question_edit_preserves_id_and_aliases_but_clears_old_checks() -> None:
    original_record = accepted_record()
    state = create_review_state(
        [original_record],
        segment_fingerprints={"01_alltypes_10": {}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )
    candidate_id = original_record["id"]

    updated = apply_review_rows(
        state,
        [review_row(candidate_id, edited_question="Who established the archive?", edit_reason="Clarity")],
    )
    artifact = Route3CandidateArtifact.from_dict(updated["candidates"][0]["artifact"])

    assert artifact.candidate_id == candidate_id
    assert len(artifact.revisions) == 2
    assert artifact.revisions[0].authoritative_question == "Who founded the archive?"
    assert artifact.current_revision.authoritative_question == "Who established the archive?"
    assert artifact.current_revision.active_aliases == ("A. Example",)
    assert artifact.current_revision.ddg == {}
    assert artifact.current_revision.second_stage == {}
    assert artifact.current_revision.status == "rerun"


def test_answer_edit_clears_aliases_and_delete_skips_rerun() -> None:
    original_record = accepted_record()
    state = create_review_state(
        [original_record],
        segment_fingerprints={"01_alltypes_10": {}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )
    candidate_id = original_record["id"]
    answer_edited = apply_review_rows(
        state,
        [review_row(candidate_id, edited_reference_answer="Grace Example")],
    )
    artifact = Route3CandidateArtifact.from_dict(answer_edited["candidates"][0]["artifact"])
    assert artifact.current_revision.active_aliases == ()

    deleted = apply_review_rows(state, [review_row(candidate_id, delete="Yes")])
    calls = []
    rerun = rerun_review_candidates(
        deleted,
        processor=lambda candidate: calls.append(candidate) or ("accepted", {}),
    )
    assert calls == []
    assert accepted_review_bundles(rerun) == []


def test_rerun_uses_fresh_checks_and_only_latest_accepted_is_exported() -> None:
    original = accepted_record()
    state = create_review_state(
        [original],
        segment_fingerprints={"01_alltypes_10": {}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )
    candidate_id = original["id"]
    edited = apply_review_rows(
        state,
        [review_row(candidate_id, edited_question="Who established the archive?")],
    )
    seen_candidates = []

    def processor(candidate):
        seen_candidates.append(deepcopy(candidate))
        assert candidate.search_verification_features == {}
        assert candidate.panel_grading_features == {}
        assert candidate.source_metadata["route3_revision_number"] == 2
        record = candidate.to_output_record("")
        record["validation"] = {"fresh": True}
        record["source_metadata"]["rule_based_qa_gate"] = {"fresh": True}
        record["search_verification_features"] = {
            "fresh": True,
            "category_hit_rates": {"overall": {"answer_hit_rate": 0.0}},
        }
        record["panel_grading_features"] = {
            "fresh": True,
            "models": deepcopy(original["panel_grading_features"]["models"]),
        }
        return "accepted", record

    rerun = rerun_review_candidates(edited, processor=processor)
    rerun = classify_review_topics(
        rerun,
        classifier=fake_topic_classifier,
        concurrency_limit=2,
    )
    artifact = Route3CandidateArtifact.from_dict(rerun["candidates"][0]["artifact"])

    assert len(seen_candidates) == 1
    assert artifact.candidate_id == candidate_id
    assert len(artifact.revisions) == 2
    assert artifact.revisions[0].ddg["passed"] is True
    assert artifact.current_revision.ddg["fresh"] is True
    assert artifact.current_revision.second_stage["fresh"] is True
    assert artifact.current_revision.status == "accepted"
    assert len(accepted_review_bundles(rerun)) == 1
    markdown = render_review_markdown(rerun, run_id="review-run")
    assert "Who established the archive?" in markdown
    assert "Human edited: `Yes`" in markdown


def test_invalid_topic_response_is_rejected_with_raw_response_retained() -> None:
    state = _create_review_state(
        [accepted_record()],
        segment_fingerprints={"01_alltypes_10": {}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )

    def invalid_classifier(_artifact, prompt):
        return {
            "text": "History.",
            "raw_text": "History.",
            "request_payload": {
                "model": "openai/gpt-4.1-mini",
                "max_tokens": 256,
            },
            "response_body": {"id": "raw-invalid", "choices": [{"message": {"content": "History."}}]},
            "raw_response_body_text": '{"id":"raw-invalid"}',
        }

    classified = classify_review_topics(
        state,
        classifier=invalid_classifier,
        concurrency_limit=1,
    )

    assert accepted_review_bundles(classified) == []
    bundle = classified["candidates"][0]
    revision = bundle["artifact"]["revisions"][-1]
    audit = bundle["topic_classifications"][0]
    assert revision["status"] == "rejected"
    assert revision["topic"] == ""
    assert audit["raw_response"]["id"] == "raw-invalid"
    assert audit["raw_response_body_text"] == '{"id":"raw-invalid"}'
    assert audit["raw_text"] == "History."
    assert audit["rejection_reason"] == "invalid_topic_classification_response"
    assert "Who founded the archive?" in audit["prompt"]
    assert "Reference answer: Ada Example" in audit["prompt"]
    assert all(topic in audit["prompt"] for topic in REVIEW_TOPICS)


def test_topic_and_human_edited_are_read_only() -> None:
    state = create_review_state(
        [accepted_record()],
        segment_fingerprints={"01_alltypes_10": {}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )
    candidate_id = state["candidates"][0]["artifact"]["id"]

    with pytest.raises(ValueError, match="topic is read-only"):
        apply_review_rows(state, [review_row(candidate_id, topic="Other")])
    with pytest.raises(ValueError, match="human_edited is read-only"):
        apply_review_rows(state, [review_row(candidate_id, human_edited="Yes")])


def test_markdown_is_sharded_by_fifty_and_model_markdown_is_quoted() -> None:
    records = [accepted_record(page_id=index) for index in range(1, 52)]
    records[0]["panel_grading_features"]["models"][0]["predicted_answer"] = (
        "# Embedded heading\n\n| A | B |\n| --- | --- |"
    )
    state = create_review_state(
        records,
        segment_fingerprints={"01_alltypes_10": {}},
        segment_artifact_roots={"01_alltypes_10": "C:/artifacts/01_alltypes_10"},
    )

    with tempfile.TemporaryDirectory() as tmpdir:
        base_path = Path(tmpdir) / "review.md"
        paths = write_review_markdown_shards(base_path, state, run_id="review-run")

        assert [path.name for path in paths] == [
            "review_1-50.md",
            "review_51-100.md",
        ]
        first = paths[0].read_text(encoding="utf-8")
        second = paths[1].read_text(encoding="utf-8")
        assert "> # Embedded heading" in first
        assert "> | A | B |" in first
        assert "\n# Embedded heading\n" not in first
        assert "## 51." in second
