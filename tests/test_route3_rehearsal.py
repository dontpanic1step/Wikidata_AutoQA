from __future__ import annotations

import csv
from copy import deepcopy
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from test_support import ROOT  # noqa: F401
from test_openrouter_batch_scripts import load_script_module
from test_route3_ddg import RecordingSearchClient, candidate as ddg_candidate, verify as verify_ddg
from test_route3_external_control import FailingSearchClient, SequenceTransport
from test_route3_openrouter import FakeTransport, RESPONSE_BODY
from test_route3_review import accepted_record, create_review_state, fake_topic_classifier
from wikidata_simpleqa.generator_validators import SearchLongtailVerifierError
from wikidata_simpleqa.public_ids import new_public_id_registry
from wikidata_simpleqa.route3_artifacts import Route3CandidateArtifact
from wikidata_simpleqa.route3_circuit import CircuitOpenError, ServiceCircuit
from wikidata_simpleqa.route3_ddg import Route3DDGVerifierResultStore
from wikidata_simpleqa.route3_external_calls import ExternalCallRecordStore
from wikidata_simpleqa.route3_external_lifecycle import (
    resolve_ambiguous_external_calls,
    scan_ambiguous_external_calls,
    write_ambiguous_external_call_reports,
)
from wikidata_simpleqa.route3_finalization import finalize_review_state, write_final_csv
from wikidata_simpleqa.route3_openrouter import (
    OpenRouterHTTPError,
    OpenRouterRawResponse,
    Route3DurableOpenRouterExecutor,
)
from wikidata_simpleqa.route3_review import (
    REVIEW_COLUMNS,
    accepted_review_bundles,
    classify_review_topics,
    apply_review_rows,
    read_review_workbook,
    render_review_markdown,
    rerun_review_candidates,
    write_review_workbook,
)
from wikidata_simpleqa.route3_run_ledger import (
    SegmentLedgerIndex,
    canonical_json_sha256,
    rebuild_derived_outputs,
    rebuild_summary_from_ledger,
)


RUN_GROUP_ID = "offline-rehearsal"
INITIAL_SEGMENT_ID = "01_alltypes_3"
TOPUP_SEGMENT_ID = "01_alltypes_3_topup_01"


def _index(root: Path, segment_id: str) -> SegmentLedgerIndex:
    segment_root = root / segment_id
    return SegmentLedgerIndex(
        allocation_dir=segment_root / "page_allocations",
        attempt_dir=segment_root / "page_attempts",
        run_group_id=RUN_GROUP_ID,
        segment_id=segment_id,
        run_group_segments_dir=root,
    )


def _attempt(
    page_id: int,
    *,
    attempt_number: int = 1,
    status: str,
    accepted_records: list[dict] | None = None,
) -> dict:
    return {
        "canonical_page_id": page_id,
        "canonical_page_url": f"https://en.wikipedia.org/?curid={page_id}",
        "attempt_number": attempt_number,
        "status": status,
        "reason": status,
        "accepted_records": accepted_records or [],
        "rejected_records": [],
        "error_details": {},
    }


def _formal_record(
    page_id: int,
    answer_type: str,
    segment_id: str,
    *,
    attempt_number: int = 1,
) -> dict:
    record = accepted_record(
        page_id=page_id,
        slot=answer_type,
        answer_type=answer_type,
        segment_id=segment_id,
    )
    record["source_metadata"]["page_attempt"] = attempt_number
    return record


def test_crash_safe_rehearsal_lifecycle(tmp_path: Path) -> None:
    archives = tmp_path / "fake_page_archives"
    archives.mkdir()
    for page_id in range(101, 106):
        (archives / f"page_{page_id}.json").write_text(
            json.dumps(
                {
                    "canonical_page_id": page_id,
                    "canonical_url": f"https://en.wikipedia.org/?curid={page_id}",
                    "title": f"Offline Page {page_id}",
                }
            ),
            encoding="utf-8",
        )

    initial = _index(tmp_path, INITIAL_SEGMENT_ID)
    for page_id in (101, 102, 103):
        archive_path = archives / f"page_{page_id}.json"
        initial.commit_allocation(
            canonical_page_id=page_id,
            page_source="cache",
            source_url=f"https://en.wikipedia.org/?curid={page_id}",
            cached_archive_path=str(archive_path),
        )

    generation_root = tmp_path / INITIAL_SEGMENT_ID / "external_calls"
    generation_request = {
        "model": "google/gemini-3-flash-preview",
        "messages": [{"role": "user", "content": "Generate all five slots."}],
    }
    generation_transport = FakeTransport(OpenRouterRawResponse(200, RESPONSE_BODY))
    generation_store = ExternalCallRecordStore(generation_root, canonical_page_id=101)
    generation_executor = Route3DurableOpenRouterExecutor(
        store=generation_store,
        transport=generation_transport,
    )
    first_generation = generation_executor.execute(
        call_key="generation",
        request_payload=generation_request,
    )

    ddg_root = tmp_path / INITIAL_SEGMENT_ID / "ddg_verifier_results"
    ddg_store = Route3DDGVerifierResultStore(ddg_root, segment_fingerprint="initial-sha")
    first_ddg = verify_ddg(ddg_store, ddg_candidate(page_id=101), RecordingSearchClient())

    resumed = _index(tmp_path, INITIAL_SEGMENT_ID)
    assert resumed.pending_primary_page_ids == [101, 102, 103]
    assert resumed.next_attempt_number(101) == 1
    assert resumed.next_attempt_number(102) == 1
    assert resumed.next_attempt_number(103) == 1

    recalled_transport = FakeTransport(AssertionError("generation was recalled"))
    resumed_generation = Route3DurableOpenRouterExecutor(
        store=generation_store,
        transport=recalled_transport,
    ).execute(call_key="generation", request_payload=generation_request)
    resumed_ddg = verify_ddg(
        ddg_store,
        ddg_candidate(page_id=101),
        RecordingSearchClient(error=AssertionError("DDG was recalled")),
    )
    assert resumed_generation == first_generation
    assert resumed_ddg == first_ddg
    assert generation_transport.calls == [generation_request]
    assert recalled_transport.calls == []

    person = _formal_record(101, "Person", INITIAL_SEGMENT_ID)
    extra_date = _formal_record(101, "Date", INITIAL_SEGMENT_ID)
    place = _formal_record(102, "Place", INITIAL_SEGMENT_ID, attempt_number=2)
    number = _formal_record(103, "Number", INITIAL_SEGMENT_ID)
    resumed.commit_attempt(
        _attempt(101, status="accepted", accepted_records=[person, extra_date])
    )
    resumed.commit_attempt(_attempt(102, status="retryable_failure"))
    assert resumed.eligible_retry_page_ids == [102]
    resumed.commit_attempt(
        _attempt(102, attempt_number=2, status="accepted", accepted_records=[place])
    )
    resumed.commit_attempt(_attempt(103, status="accepted", accepted_records=[number]))

    ambiguous_request = {
        "model": "google/gemini-3-flash-preview",
        "messages": [{"role": "user", "content": "Grade the Other slot."}],
    }
    ambiguous_store = ExternalCallRecordStore(generation_root, canonical_page_id=103)
    ambiguous_store.commit(
        call_key="second_stage_answer_other_gemini",
        record_kind="intent",
        request_hash=canonical_json_sha256(ambiguous_request),
        payload={"request_payload": ambiguous_request},
    )
    ambiguous_json = tmp_path / "ambiguous_external_calls.json"
    ambiguous_md = tmp_path / "ambiguous_external_calls.md"
    ambiguous_rows = write_ambiguous_external_call_reports(
        generation_root,
        json_path=ambiguous_json,
        markdown_path=ambiguous_md,
    )
    assert len(ambiguous_rows) == 1
    assert "Unresolved calls: 1" in ambiguous_md.read_text(encoding="utf-8")
    resolved = resolve_ambiguous_external_calls(generation_root, action="abandon")
    assert resolved[0]["resolution_action"] == "abandon"
    assert scan_ambiguous_external_calls(generation_root) == []

    openrouter_circuit = ServiceCircuit("openrouter")
    failing_transport = SequenceTransport(
        [
            OpenRouterHTTPError(status_code=429, body_text="rate"),
            OpenRouterHTTPError(status_code=500, body_text="server"),
            OpenRouterHTTPError(status_code=408, body_text="timeout"),
        ]
    )
    for page_id in (201, 202, 203):
        executor = Route3DurableOpenRouterExecutor(
            store=ExternalCallRecordStore(tmp_path / "circuit_calls", canonical_page_id=page_id),
            transport=failing_transport,
            circuit=openrouter_circuit,
        )
        with pytest.raises(Exception):
            executor.execute(
                call_key="generation",
                request_payload={"model": "test/model", "messages": [page_id]},
            )
    assert openrouter_circuit.is_open
    blocked_transport = SequenceTransport([])
    with pytest.raises(CircuitOpenError):
        Route3DurableOpenRouterExecutor(
            store=ExternalCallRecordStore(tmp_path / "circuit_calls", canonical_page_id=204),
            transport=blocked_transport,
            circuit=openrouter_circuit,
        ).execute(
            call_key="generation",
            request_payload={"model": "test/model", "messages": [204]},
        )
    assert failing_transport.calls == 3
    assert blocked_transport.calls == 0
    resumed_transport = SequenceTransport([OpenRouterRawResponse(200, RESPONSE_BODY)])
    Route3DurableOpenRouterExecutor(
        store=ExternalCallRecordStore(tmp_path / "circuit_calls", canonical_page_id=204),
        transport=resumed_transport,
        circuit=ServiceCircuit("openrouter"),
    ).execute(
        call_key="generation",
        request_payload={"model": "test/model", "messages": [204]},
    )
    assert resumed_transport.calls == 1

    ddg_circuit = ServiceCircuit("duckduckgo")
    circuit_ddg_store = Route3DDGVerifierResultStore(
        tmp_path / "circuit_ddg",
        segment_fingerprint="circuit-sha",
        circuit=ddg_circuit,
    )
    for page_id in (201, 202, 203):
        with pytest.raises(SearchLongtailVerifierError):
            verify_ddg(circuit_ddg_store, ddg_candidate(page_id=page_id), FailingSearchClient())
    assert ddg_circuit.is_open
    blocked_search = FailingSearchClient()
    with pytest.raises(CircuitOpenError):
        verify_ddg(circuit_ddg_store, ddg_candidate(page_id=204), blocked_search)
    assert blocked_search.calls == 0
    recovery_search = RecordingSearchClient()
    recovery_ddg_store = Route3DDGVerifierResultStore(
        tmp_path / "circuit_ddg",
        segment_fingerprint="circuit-sha",
        circuit=ServiceCircuit("duckduckgo"),
    )
    assert verify_ddg(recovery_ddg_store, ddg_candidate(page_id=204), recovery_search)[0]
    assert len(recovery_search.calls) == 2

    topup = _index(tmp_path, TOPUP_SEGMENT_ID)
    with pytest.raises(ValueError, match="already allocated page 101"):
        topup.commit_allocation(canonical_page_id=101, page_source="cache")
    for page_id in (104, 105):
        topup.commit_allocation(
            canonical_page_id=page_id,
            page_source="fresh",
            source_url=f"https://en.wikipedia.org/?curid={page_id}",
            cached_archive_path=str(archives / f"page_{page_id}.json"),
        )
    date = _formal_record(104, "Date", TOPUP_SEGMENT_ID)
    other = _formal_record(105, "Other", TOPUP_SEGMENT_ID)
    topup.commit_attempt(_attempt(104, status="accepted", accepted_records=[date]))
    topup.commit_attempt(_attempt(105, status="accepted", accepted_records=[other]))

    assert initial.primary_page_ids.isdisjoint(topup.primary_page_ids)
    assert topup.run_group_page_ids == {101, 102, 103, 104, 105}
    assert len(list(archives.glob("page_*.json"))) == 5
    for allocation in resumed.allocations + topup.allocations:
        assert Path(allocation["cached_archive_path"]).exists()
    for index in (resumed, topup):
        attempts_by_page: dict[int, list[int]] = {}
        for attempt in index.attempts:
            attempts_by_page.setdefault(int(attempt["canonical_page_id"]), []).append(
                int(attempt["attempt_number"])
            )
        assert all(numbers in ([1], [1, 2]) for numbers in attempts_by_page.values())

    accepted_records: list[dict] = []
    segment_fingerprints = {}
    segment_roots = {}
    for index, fingerprint in ((resumed, "initial-sha"), (topup, "topup-sha")):
        segment_root = index.allocation_dir.parent
        accepted_path = segment_root / "accepted.jsonl"
        rejected_path = segment_root / "rejected.jsonl"
        summary_path = segment_root / "summary.json"
        accepted, rejected, _retry_pending = rebuild_derived_outputs(
            index,
            accepted_path=accepted_path,
            rejected_path=rejected_path,
        )
        summary = rebuild_summary_from_ledger(summary_path, index)
        expected_accepted = accepted_path.read_text(encoding="utf-8")
        expected_rejected = rejected_path.read_text(encoding="utf-8")
        accepted_path.unlink()
        rejected_path.unlink()
        summary_path.unlink()
        rebuilt_accepted, rebuilt_rejected, _ = rebuild_derived_outputs(
            index,
            accepted_path=accepted_path,
            rejected_path=rejected_path,
        )
        rebuilt_summary = rebuild_summary_from_ledger(summary_path, index)
        assert rebuilt_accepted == accepted
        assert rebuilt_rejected == rejected
        assert accepted_path.read_text(encoding="utf-8") == expected_accepted
        assert rejected_path.read_text(encoding="utf-8") == expected_rejected
        assert rebuilt_summary == summary
        accepted_records.extend(rebuilt_accepted)
        segment_fingerprints[index.segment_id] = {"sha256": fingerprint}
        segment_roots[index.segment_id] = str(segment_root)

    review_state = create_review_state(
        accepted_records,
        segment_fingerprints=segment_fingerprints,
        segment_artifact_roots=segment_roots,
    )
    review_md = tmp_path / "review.md"
    review_xlsx = tmp_path / "review.xlsx"
    review_md.write_text(render_review_markdown(review_state, run_id=RUN_GROUP_ID), encoding="utf-8")
    write_review_workbook(review_xlsx, review_state)
    workbook = load_workbook(review_xlsx)
    worksheet = workbook["review"]
    assert tuple(cell.value for cell in worksheet[1]) == REVIEW_COLUMNS
    ids_by_page = {
        int(bundle["artifact"]["provenance"]["canonical_page_id"]): str(bundle["artifact"]["id"])
        for bundle in review_state["candidates"]
    }
    extra_date_id = extra_date["id"]
    place_id = place["id"]
    for row_number in range(2, worksheet.max_row + 1):
        candidate_id = str(worksheet.cell(row=row_number, column=1).value)
        if candidate_id == extra_date_id:
            worksheet.cell(row=row_number, column=7).value = "Yes"
        if candidate_id == place_id:
            worksheet.cell(row=row_number, column=8).value = "Who established the archive?"
            worksheet.cell(row=row_number, column=10).value = "Clarity"
    workbook.save(review_xlsx)
    workbook.close()

    applied = apply_review_rows(review_state, read_review_workbook(review_xlsx))
    revalidation_calls = []

    def processor(candidate):
        revalidation_calls.append(deepcopy(candidate))
        assert candidate.search_verification_features == {}
        assert candidate.panel_grading_features == {}
        record = candidate.to_output_record("")
        record["validation"] = {"fresh": True}
        record["source_metadata"]["rule_based_qa_gate"] = {"fresh": True}
        record["search_verification_features"] = {
            "fresh": True,
            "category_hit_rates": {"overall": {"answer_hit_rate": 0.0}},
        }
        record["panel_grading_features"] = {
            "fresh": True,
            "models": deepcopy(place["panel_grading_features"]["models"]),
        }
        return "accepted", record

    revalidated = rerun_review_candidates(applied, processor=processor)
    revalidated = classify_review_topics(
        revalidated,
        classifier=fake_topic_classifier,
        concurrency_limit=2,
    )
    assert len(revalidation_calls) == 1
    assert revalidation_calls[0].final_question == "Who established the archive?"
    active_ids = {str(bundle["artifact"]["id"]) for bundle in accepted_review_bundles(revalidated)}
    assert extra_date_id not in active_ids
    assert place_id in active_ids

    final_xlsx = tmp_path / "review_revalidated.xlsx"
    write_review_workbook(final_xlsx, revalidated)
    final_rows = read_review_workbook(final_xlsx)
    final_result = finalize_review_state(
        revalidated,
        review_rows=final_rows,
        public_id_registry=new_public_id_registry(),
    )
    final_csv = tmp_path / "final.csv"
    write_final_csv(final_csv, final_result["records"])
    with final_csv.open(encoding="utf-8", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    assert len(csv_rows) == 5
    assert {row["answer_type"] for row in csv_rows} == {"Person", "Place", "Number", "Date", "Other"}
    assert any(row["problem"] == "Who established the archive?" for row in csv_rows)
    assert all(row["id"] != extra_date_id for row in csv_rows)
    assert set(ids_by_page) == {101, 102, 103, 104, 105}

    grader_module = load_script_module(
        "judge_openrouter_batch_predictions_rehearsal_snapshot",
        "scripts/judge_openrouter_batch_predictions.py",
    )
    prompt = grader_module.GRADER_TEMPLATE.format(
        question="Which city hosted the example event?",
        target="Example City",
        predicted_answer="It was held in Example City.",
    )
    assert hashlib.sha256(grader_module.GRADER_TEMPLATE.encode("utf-8")).hexdigest() == (
        "84c004ec4fcf8f0703bb0d734544036a72e847bfa7116429e5aee5e53ccc8cf3"
    )
    assert hashlib.sha256(prompt.encode("utf-8")).hexdigest() == (
        "90ac762a5286b75464a16f6132898cd7a4a7c9f726ee64ce21034a905cc328fe"
    )
